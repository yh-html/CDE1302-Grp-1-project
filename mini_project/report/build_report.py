"""
Builds the SKELETON of the CDE1302 mini-project group report as a Word file.

Format rules from the brief: at most 10 pages including cover and appendix, single
column, single spacing, Times New Roman at least 12 pt, GitHub link, concise
integrated-system description.

Text in [[double brackets]] is a placeholder the group must replace; it is
highlighted yellow.  Everything else is a DRAFT written from the code in this
repository: check each statement against your own drone before submitting.

Run once to create the skeleton.  After the group starts editing the .docx, do NOT
re-run this (it would overwrite the edits).
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "CDE1302_Mini_Project_Report.docx"
FONT = "Times New Roman"
SIZE = 12

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)           # A4
sec.left_margin = sec.right_margin = Inches(1.0)
sec.top_margin = sec.bottom_margin = Inches(1.0)


def _set_font(style, size=SIZE, bold=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    # theme font attributes override the names above (headings would fall back to Calibri)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


normal = doc.styles["Normal"]
_set_font(normal)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.space_before = Pt(0)
for name, size in (("Heading 1", 13), ("Heading 2", 12)):
    st = doc.styles[name]
    _set_font(st, size, bold=True)
    st.font.color.rgb = None
    st.paragraph_format.space_before = Pt(8)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    st.paragraph_format.keep_with_next = True
for name in ("List Bullet", "Caption"):
    _set_font(doc.styles[name])
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)       # Caption defaults to blue
    doc.styles[name].font.bold = False
    doc.styles[name].paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

TODO = re.compile(r"(\[\[.*?\]\])")


def runs(p, text, bold=False, italic=False):
    for part in TODO.split(text):
        if not part:
            continue
        r = p.add_run(part.replace("[[", "[").replace("]]", "]") if TODO.fullmatch(part) else part)
        r.bold, r.italic = bold, italic
        r.font.name, r.font.size = FONT, Pt(SIZE)
        if TODO.fullmatch(part):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def para(text, bold=False, italic=False, align=None, after=None):
    p = doc.add_paragraph()
    runs(p, text, bold, italic)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    runs(p, text)
    p.paragraph_format.space_after = Pt(1)
    return p


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def caption(text):
    p = doc.add_paragraph(style="Caption")
    runs(p, text, italic=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def table(header, rows, widths=None, small=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        runs(c.paragraphs[0], h, bold=True)
        shade(c, "D9D9D9")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            runs(cells[i].paragraphs[0], str(v))
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def placeholder_box(text, height_in=1.6):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    c.text = ""
    shade(c, "FFF2CC")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    runs(p, text)
    tr = t.rows[0]._tr
    trpr = tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(height_in * 1440)))
    h.set(qn("w:hRule"), "atLeast")
    trpr.append(h)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ---------------------------------------------------------------------------
# Title block (compact: the cover page counts toward the 10-page limit)
# ---------------------------------------------------------------------------
p = para("CDE1302 Mini Project: Autonomous UAV Exploration", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p.runs[0].font.size = Pt(16)
para("Group Report", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Lab group: [[group / lab session]]    AY: [[2026/27]]    Semester: [[1]]    "
     "Demonstration date: [[date]]", align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Members: [[Name, Matric No.]]; [[Name, Matric No.]]; [[Name, Matric No.]]; [[Name, Matric No.]]",
     align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Source code: [[GitHub repository link (the repo that holds the final code)]]",
     align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

# ---------------------------------------------------------------------------
h1("1. Introduction and mission")
para("The UAV must take off on a command from the ground laptop, search a bounded area for AprilTags at unknown "
     "positions, record each tag's ID, estimated position in the UAV's local frame, relative pose and timestamp, "
     "return to the start and land, with no manual input after takeoff. We describe one integrated system: a "
     "flight controller for stabilisation, a companion computer for tag detection and command execution, and a "
     "laptop program that plans the search, localises tags, logs events and checks safety. "
     "[[One or two sentences on your result: tags found / total, mean localisation error, mission time.]]")

h1("2. System architecture and data/control flow")
para("[[Hardware: state your final configuration, for example frame and camera mount, motors, props, battery, "
     "and any modifications with the reason for each.]] The vehicle is a PX4 quadrotor (MicoAir743 v2 flight "
     "controller) with an optical-flow sensor (MTF-02P) for local position, and an ESP32-S3 companion with a camera "
     "that detects AprilTags (tag16h5 family) on board.")
doc.add_picture(str(HERE / "fig_architecture.png"), width=Inches(5.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption("Figure 1. System architecture and data flow (drawn from the code: ports and rates as implemented).")
para("Control flow: the laptop sends CMD_START; the ESP32 switches PX4 to offboard mode, arms and climbs to 0.5 m. "
     "The laptop then sends CMD_GOTO waypoints; the ESP32 streams position setpoints to PX4 at 50 Hz, and PX4's "
     "position controller flies the drone, so the flight speed is set by PX4 parameters, not by the laptop. "
     "Data flow: telemetry (local x, y, heading, navigation state) returns to the laptop at 10 Hz; tag detections "
     "(ID and camera-frame position) arrive separately on a second stream at about the camera processing rate. "
     "The camera path is read-only with respect to flight: detections never change the flight commands. "
     "PX4 records its own flight log (.ulg), which is used after the flight (Section 7).")

h1("3. Exploration and navigation strategy")
para("The laptop flies a boundary-limited lawn-mower sweep. We chose it because the brief only requires staying "
     "inside the boundaries and allows waypoint navigation, and because a deterministic sweep is simple, "
     "predictable and easy to verify. The drone has no obstacle sensing in its current firmware (the time-of-flight "
     "avoidance stack is disabled), so the method does not rely on it.")
bullet("Planning: the search area (announced before the demonstration) is shrunk by an edge margin; lanes run "
       "along a configured axis (forward by default); the lane order and direction are chosen to minimise the flown "
       "distance including the return.")
bullet("Lane spacing: derived from the camera, as twice the sideways distance at which a floor tag can be seen "
       "(tilt, field of view, altitude, range limit, minimum viewing angle), reduced by a safety factor of 0.8. "
       "[[Your final lane spacing and how you checked it with a real tag.]]")
bullet("Execution: each waypoint is sent as CMD_GOTO and accepted within 0.25 m (the firmware arrival radius); "
       "a missed waypoint triggers a hold and continues; three consecutive misses end the search early.")
bullet("Return and landing: after the last lane the drone returns to the takeoff point and the landing command is "
       "sent. Landing is commanded, then a fixed wait follows, because telemetry has no touchdown signal.")
bullet("Safety: the program lands if telemetry is stale (3 s), if the drone leaves the search area by more than "
       "0.5 m, on operator interrupt, or on any error. It stops searching early when the time limit is reached or "
       "when all expected tags are found (optional).")
para("Alternatives considered: [[the swarm stack's gap-based exploration (needs the time-of-flight sensors, "
     "disabled in our firmware), trajectory tracking with an uploaded minimum-snap path, a random walk. State why "
     "you did or did not use each.]]")

h1("4. AprilTag detection and target localisation")
para("Detection runs on the ESP32 (tag16h5 family). A detection is accepted only if it has at most one corrected bit "
     "error, a decision margin above 55 and a pose reprojection error below 0.5, which are the firmware's quality "
     "gates. For every accepted tag the ESP32 reports the tag's position in the camera frame (X right, Y down, "
     "Z forward). The laptop converts it to the UAV's local frame:")
para("p_body = R_y(θ)·p_cam + t_cam,        p_local = p_drone + R_z(ψ)·[p_body,x ; p_body,y]", italic=True,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("where θ is the camera's nose-down tilt, t_cam is the lens offset from the drone body origin, p_drone is the "
     "drone position from telemetry interpolated to the frame time, and ψ is the heading. Repeated sightings of the "
     "same ID are fused by a weighted mean with weight 1/range², so nearer (more accurate) sightings count more. "
     "The tag's relative pose is recorded as the camera-frame translation: the detection stream does not carry the "
     "tag's rotation. Every sighting, accepted or not, is written to a log with all intermediate values.")
para("Assumptions and error sources: [[camera tilt and lens offset as measured on the assembled drone (state the "
     "values and how you measured them); the drone's roll and pitch are not in the telemetry and are ignored "
     "(Section 7 quantifies this from the flight log); detection latency; tags assumed on the floor.]]")

h1("5. Important parameters and implementation decisions")
para("All tunable values are in one configuration file (tag_mission.yaml), as the brief asks. Values marked * are "
     "assumptions that must be replaced by your measurements.")
table(["Parameter", "Value", "Basis"], [
    ["Search area (local frame)", "[[from examiner]]", "announced before the demo"],
    ["Cruise altitude", "0.5 m", "firmware constant CRUISE_ALT_M"],
    ["Cruise speed (timeouts only)", "[[measured]] m/s *", "PX4 MPC_XY_VEL_MAX; measure from the .ulg"],
    ["Lane spacing", "[[value]] m *", "camera swath x 0.8"],
    ["Camera tilt / lens offset", "[[45°]] / [[x, y, z]] mm *", "measure on the drone"],
    ["Camera field of view", "91° x 82°", "firmware calibration"],
    ["Max tag range / min viewing angle", "[[3 m]] / [[30°]] *", "test with a real tag"],
    ["Waypoint arrival radius", "0.25 m", "firmware"],
    ["Safety limits", "stale 3 s, fence +0.5 m", "configuration"],
], widths=[2.3, 1.7, 2.4])
para("Key decisions: [[list 3 to 5, each with the reason: for example the lawn-mower strategy; using the "
     "stock frame with a rigid camera mount; position-setpoint waypoints instead of trajectory upload; "
     "laptop-side localisation so the geometry can be changed without reflashing.]]")

h1("6. Experimental results and performance observations")
para("[[Describe the test setup: area size, number and placement of tags, how true tag positions were measured "
     "(tape measure, reference marks), lighting, number of trials.]]")
table(["Trial", "Tags found / placed", "Mean error (m)", "Max error (m)", "Mission time (s)", "Notes"], [
    ["1", "[[ ]]", "[[ ]]", "[[ ]]", "[[ ]]", "[[ ]]"],
    ["2", "[[ ]]", "[[ ]]", "[[ ]]", "[[ ]]", "[[ ]]"],
    ["3", "[[ ]]", "[[ ]]", "[[ ]]", "[[ ]]", "[[ ]]"],
], widths=[0.5, 1.2, 1.0, 1.0, 1.1, 1.6])
para("[[Observations: detection reliability by range and angle, localisation error against range, how closely the "
     "drone returned to the start, measured flight speed against the configured value, any failures and what "
     "caused them. Report what happened, including failed trials.]]")
para("Software verification: the laptop code has automated tests, including full missions in a simulator and over "
     "real UDP sockets against a fake drone. These test the logic, not real-world accuracy, because the simulator "
     "uses the same camera geometry as the localiser. [[Update the test count from the final run.]]")

h1("7. AI-assisted flight-data analysis and visualisation")
para("After each flight the PX4 .ulg log was downloaded and combined with the mission log by a Python script. It "
     "aligns the two clocks (PX4 time since boot against the laptop wall clock) by matching the drone path, then "
     "plots the full trajectory, the start and landing positions, and the estimated tag positions with their IDs, "
     "and tabulates path length, return error and localisation error against the measured tag positions.")
placeholder_box("[[Figure 2. Flight overview from your real flight: trajectory, start, landing, tag IDs and "
                "positions (output of analyze_flight.py, flight_overview.png).]]", 2.6)
para("The script also re-localises every sighting using the logged roll and pitch, to measure the error from "
     "ignoring tilt in flight. [[Report the shift per tag from your flight.]]")
h2("How AI was used, its limits, and how results were verified")
para("[[Tools and what each was used for (code, plots, analysis). Keep this factual and specific to your group.]] "
     "Examples of AI errors that verification caught during development (replace or extend with your own):")
bullet("An AI assistant stated that the flight speed was set by a firmware constant; reading the code showed the "
       "constant is unused, and the speed is a PX4 parameter. It was caught by reading the source and the repository's "
       "own test guide, and the measured speed from the log now replaces the assumption. [[Confirm with your log.]]")
bullet("A first lane-spacing estimate treated any tag within camera range as detectable and suggested lanes 3 m "
       "apart; adding a minimum viewing angle (a floor tag seen edge-on will not decode) reduced it to about 1 m. "
       "[[Confirm with a real tag test.]]")
para("Verification: unit tests on known geometry; synthetic logs in the PX4 log format; comparison of estimated "
     "against tape-measured tag positions; checking the flight-path plot against the observed flight. "
     "[[Add what you actually did, and what remained unverified.]]")

h1("8. Limitations, failure cases and possible improvements")
bullet("No obstacle avoidance (sensors disabled in the firmware); the sweep assumes a clear, bounded area.")
bullet("Localisation depends on the measured camera mount; the telemetry has no roll or pitch, and detection "
       "latency shifts the pose used. [[Quantified error from your results.]]")
bullet("Landing is a command plus a wait; no touchdown signal in telemetry. Auto-landing reliability: "
       "[[describe what you observed.]]")
bullet("Completion time is limited by the PX4 speed limit; raising it increases motion blur and pose lag. "
       "[[What you tried.]]")
bullet("The tag packet carries translation only, so the tag's orientation is not recorded.")
para("Improvements: [[for example adaptive search that revisits gaps, using tag height information, adding "
     "roll and pitch to the telemetry, speed tuning with detection latency compensation, automatic recovery.]]")

h1("Appendix A. Team contributions and reproduction")
table(["Member", "Contribution"], [
    ["[[Name]]", "[[e.g. airframe and camera mount, hardware assembly and tuning]]"],
    ["[[Name]]", "[[e.g. mission software, localisation, tests]]"],
    ["[[Name]]", "[[e.g. flight testing and data analysis]]"],
    ["[[Name]]", "[[e.g. report, verification, demonstration]]"],
], widths=[1.6, 4.8])
para("Run: python laptop/run_tag_mission.py --drone-id [[ID]]   (dry run: --simulate; plan only: --plan-only). "
     "Analyse: python laptop/analyze_flight.py --ulg [[file.ulg]] --run [[run folder]]. "
     "Repository: [[link]]. Source code is excluded from the page limit; do not paste it here.")

doc.save(OUT)
print("wrote", OUT)
